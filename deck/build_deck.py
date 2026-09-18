#!/usr/bin/env python3
"""build_deck.py — render the Java Memory Management / Garbage Collector deck.

Pipeline
--------
1. build one HTML file per slide (1920x1080 design canvas)
2. render each slide to PNG with headless Chromium at 2x (3840x2160)
3. assemble a real 16:9 PowerPoint (.pptx) with full-bleed slides + notes
4. export a print-ready PDF handout

Run:  python3 build_deck.py            (everything)
      python3 build_deck.py --html     (only write the HTML)
      python3 build_deck.py --png      (only render the PNGs)
"""
import os
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out")
SLIDE_DIR = os.path.join(OUT, "html")
PNG_DIR = os.path.join(OUT, "png")
PROJECT = os.path.dirname(HERE)

DECK_FILE = os.path.join(PROJECT, "Java_Memory_Management_GC.pptx")
PDF_FILE = os.path.join(PROJECT, "Java_Memory_Management_GC.pdf")

CSS = ["css/base.css", "css/cards.css", "css/editor.css", "css/diagrams.css"]

PAGE = """<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<title>{title}</title>
{links}
<style>
  html, body {{ margin:0; padding:0; background:#05070C; }}
  .slide {{ margin:0; }}
</style>
</head><body>
<div class="slide">
  <div class="vignette"></div>
  {chrome}
  <div class="stage" style="{stage_style}">{body}</div>
</div>
</body></html>
"""

PRINT_PAGE = """<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"><title>Java Memory Management — GC</title>
{links}
<style>
  @page {{ size: 13.3333in 7.5in; margin: 0; }}
  html, body {{ margin:0; padding:0; background:#05070C; }}
  .slide {{ page-break-after: always; break-after: page; margin:0; }}
  .slide:last-child {{ page-break-after: auto; }}
</style>
</head><body>
{slides}
</body></html>
"""


def load_slides(only=None):
    """Collect the SLIDE dict from slide01.py ... slide15.py in order."""
    sys.path.insert(0, HERE)
    slides = []
    for i in (only or range(1, 16)):
        try:
            mod = __import__(f"slide{i:02d}")
        except ModuleNotFoundError:
            continue
        slides.append(mod.SLIDE)
    slides.sort(key=lambda s: s["n"])
    return slides


def write_html(slides):
    os.makedirs(SLIDE_DIR, exist_ok=True)
    links = "\n".join(
        f'<link rel="stylesheet" href="{os.path.join("..", "..", c)}">' for c in CSS
    )
    files = []
    bodies = []
    for s in slides:
        title = f'{s["n"]:02d} — {s["kicker"]}'
        html = PAGE.format(title=title, links=links, chrome=s["chrome"],
                           stage_style=s.get("stage_style", ""), body=s["body"])
        path = os.path.join(SLIDE_DIR, f"slide-{s['n']:02d}.html")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(html)
        files.append(path)
        bodies.append(f'<div class="slide"><div class="vignette"></div>{s["chrome"]}'
                      f'<div class="stage" style="{s.get("stage_style", "")}">{s["body"]}</div></div>')
    with open(os.path.join(OUT, "print.html"), "w", encoding="utf-8") as fh:
        fh.write(PRINT_PAGE.format(links=links, slides="\n".join(bodies)))
    return files


def render_pngs(files):
    os.makedirs(PNG_DIR, exist_ok=True)
    pngs = []
    for path in files:
        name = os.path.splitext(os.path.basename(path))[0] + ".png"
        png = os.path.join(PNG_DIR, name)
        cmd = [
            "chromium", "--headless=new", "--no-sandbox", "--disable-gpu",
            "--hide-scrollbars", "--allow-file-access-from-files",
            "--force-device-scale-factor=2", "--window-size=1920,1080",
            "--virtual-time-budget=4000", f"--screenshot={png}",
            "file://" + path,
        ]
        subprocess.run(cmd, check=True, capture_output=True)
        pngs.append(png)
        print(f"  rendered {name}")
    return pngs


def render_pdf():
    print_path = os.path.join(OUT, "print.html")
    cmd = [
        "chromium", "--headless=new", "--no-sandbox", "--disable-gpu",
        "--allow-file-access-from-files", "--no-pdf-header-footer",
        "--virtual-time-budget=6000", f"--print-to-pdf={PDF_FILE}",
        "file://" + print_path,
    ]
    subprocess.run(cmd, check=True, capture_output=True)
    print(f"  pdf -> {PDF_FILE}")


def build_pptx(slides, pngs):
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13.333 in
    prs.slide_height = Emu(6858000)   # 7.5 in
    blank = prs.slide_layouts[6]
    for s, png in zip(slides, pngs):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(png, 0, 0, width=prs.slide_width,
                                 height=prs.slide_height)
        notes = slide.notes_slide.notes_text_frame
        notes.text = s["notes"].strip()
    prs.save(DECK_FILE)
    print(f"  pptx -> {DECK_FILE}")


def main():
    args = sys.argv[1:]
    only = None
    for a in args:
        if a.startswith("--only"):
            only = [int(x) for x in a.split("=", 1)[1].split(",")]
    slides = load_slides(only)
    files = write_html(slides)
    print(f"wrote {len(files)} slide HTML files")
    if "--html" in args:
        return
    pngs = render_pngs(files)
    if "--png" in args:
        return
    render_pdf()
    build_pptx(slides, pngs)
    print(f"done — {len(slides)} slides")


if __name__ == "__main__":
    main()